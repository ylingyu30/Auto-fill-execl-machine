import json
import random
import os
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

def load_wordbank(file_path, encoding='utf-8'):
    with open(file_path, 'r', encoding=encoding) as f:
        words = [line.strip() for line in f if line.strip()]
    return words

def load_row_wordbank(file_path, delimiter=',', encoding='utf-8'):
    rows = []
    with open(file_path, 'r', encoding=encoding) as f:
        for line in f:
            line = line.strip()
            if not line:
                continue
            fields = line.split(delimiter)
            rows.append(fields)
    return rows

def parse_cell_range(range_str):
    range_str = range_str.upper().strip()
    if ':' in range_str:
        start, end = range_str.split(':')
        start_col = ''.join([c for c in start if c.isalpha()])
        start_row = int(''.join([c for c in start if c.isdigit()]))
        end_col = ''.join([c for c in end if c.isalpha()])
        end_row = int(''.join([c for c in end if c.isdigit()]))
        if start_row != end_row:
            raise ValueError(f"范围 {range_str} 跨越了多行，当前仅支持单行填充")
        return start_col, start_row, end_col
    else:
        col_part = ''.join([c for c in range_str if c.isalpha()])
        row_part = ''.join([c for c in range_str if c.isdigit()])
        if not row_part:
            raise ValueError(f"无效的单元格表示：{range_str}")
        row = int(row_part)
        if len(col_part) == 1:
            return col_part, row, None
        else:
            start_col = col_part[0]
            end_col = col_part[-1]
            return start_col, row, end_col

def get_column_index(col_letter):
    index = 0
    for ch in col_letter:
        index = index * 26 + (ord(ch) - ord('A') + 1)
    return index

def is_invalid_score(cell_value, invalid_markers=None):
    if cell_value is None:
        return True
    if invalid_markers:
        if isinstance(cell_value, str):
            for marker in invalid_markers:
                if marker in cell_value:
                    return True
    try:
        float(cell_value)
        return False
    except (ValueError, TypeError):
        return True

def select_wordbank_by_score(score, thresholds, wordbanks):
    if score is None:
        return wordbanks[0]
    for i, th in enumerate(thresholds):
        if score >= th:
            return wordbanks[i]
    return wordbanks[-1]

def fill_cell_with_condition(worksheet, target_cell, score_cell, thresholds, wordbanks,
                             skip_on_invalid=True, invalid_markers=None):
    cell_value = worksheet[score_cell].value
    invalid = is_invalid_score(cell_value, invalid_markers)

    if invalid:
        if skip_on_invalid:
            print(f"跳过填充 {target_cell}：分数单元格 {score_cell} 无效 (值: {cell_value})")
            return
        else:
            score = None
            print(f"警告：{score_cell} 无效，随机使用第一个词库填充 {target_cell}")
    else:
        score = float(cell_value)

    chosen_wordbank = select_wordbank_by_score(score, thresholds, wordbanks)
    words = load_wordbank(chosen_wordbank)
    if not words:
        print(f"警告：词库 {chosen_wordbank} 为空，跳过 {target_cell}")
        return
    chosen = random.choice(words)
    worksheet[target_cell] = chosen
    print(f"已填入 {target_cell} <- 分数 {score} -> 词库 {chosen_wordbank} -> {chosen}")

def fill_cell_simple(worksheet, cell_ref, wordbank_file):
    words = load_wordbank(wordbank_file)
    if not words:
        print(f"警告：词库 {wordbank_file} 为空，跳过 {cell_ref}")
        return
    chosen = random.choice(words)
    worksheet[cell_ref] = chosen
    print(f"已填入 {cell_ref} <- {chosen}")

def fill_row_range(worksheet, start_col, row, end_col, wordbank_file, delimiter=','):
    rows_data = load_row_wordbank(wordbank_file, delimiter)
    if not rows_data:
        print(f"警告：行词库 {wordbank_file} 为空，跳过行 {row}")
        return
    chosen_row = random.choice(rows_data)
    start_idx = get_column_index(start_col)
    end_idx = get_column_index(end_col)
    expected_fields = end_idx - start_idx + 1
    if len(chosen_row) < expected_fields:
        print(f"警告：词库行字段数({len(chosen_row)})不足预期({expected_fields})，不足部分留空")
        chosen_row += [''] * (expected_fields - len(chosen_row))
    for offset, field in enumerate(chosen_row[:expected_fields]):
        target_col_idx = start_idx + offset
        target_col_letter = get_column_letter(target_col_idx)
        cell = f"{target_col_letter}{row}"
        worksheet[cell] = field
    print(f"已填入行 {row}, 列 {start_col}-{end_col} <- {chosen_row[:expected_fields]}")

def process_one_excel(input_path, output_path, subjects_config):
    try:
        wb = load_workbook(input_path)
    except Exception as e:
        print(f"无法打开文件 {input_path}: {e}")
        return False

    for subject in subjects_config:
        sheet_name = subject.get('sheet_name')
        rules = subject.get('fill_rules', {})
        if not sheet_name:
            print(f"警告：学科配置缺少 sheet_name，跳过")
            continue
        if sheet_name not in wb.sheetnames:
            print(f"警告：文件 {input_path} 中不存在工作表 '{sheet_name}'，跳过该学科")
            continue
        ws = wb[sheet_name]
        print(f"  处理学科: {subject.get('name', sheet_name)} (工作表: {sheet_name})")

        for target, rule in rules.items():
            start_col, row, end_col = parse_cell_range(target)

            if end_col is not None:
                if isinstance(rule, dict) and 'file' in rule:
                    wordbank_file = rule.get('file')
                    delimiter = rule.get('delimiter', ',')
                    if not os.path.exists(wordbank_file):
                        print(f"错误：词库文件 {wordbank_file} 不存在，跳过 {target}")
                        continue
                    fill_row_range(ws, start_col, row, end_col, wordbank_file, delimiter)
                else:
                    print(f"错误：整行填充 {target} 需配置为 {{'file': '...', 'delimiter': '...'}}")
                continue

            if isinstance(rule, dict) and 'score_cell' in rule:
                score_cell = rule['score_cell']
                thresholds = rule.get('thresholds', [75, 50])
                wordbanks = rule.get('wordbanks', [])
                if len(wordbanks) != len(thresholds) + 1:
                    print(f"错误：{target} 的词库数量应为 {len(thresholds)+1} 个，实际 {len(wordbanks)}")
                    continue
                skip_on_invalid = rule.get('skip_on_invalid', True)
                invalid_markers = rule.get('invalid_markers', [])
                fill_cell_with_condition(ws, target, score_cell, thresholds, wordbanks,
                                         skip_on_invalid, invalid_markers)
            elif isinstance(rule, str):
                fill_cell_simple(ws, target, rule)
            else:
                print(f"错误：单元格 {target} 的配置无效，应为字符串或条件字典")

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    wb.save(output_path)
    print(f"已保存: {output_path}")
    return True

def main(config_file='fill_config.json'):
    if not os.path.exists(config_file):
        print(f"配置文件 {config_file} 不存在，请创建。示例：")
        example = {
            "input_dir": "data",
            "output_dir": "output",
            "subjects": [
                {
                    "name": "英语",
                    "sheet_name": "英语",
                    "fill_rules": {
                        "B5": {
                            "score_cell": "E3",
                            "thresholds": [75, 50],
                            "wordbanks": ["wordbanks/英语/B5good.txt", "wordbanks/英语/B5med.txt", "wordbanks/英语/B5low.txt"]
                        }
                    }
                }
            ]
        }
        print(json.dumps(example, indent=2, ensure_ascii=False))
        return

    with open(config_file, 'r', encoding='utf-8') as f:
        config = json.load(f)

    input_dir = config.get('input_dir')
    output_dir = config.get('output_dir')
    subjects = config.get('subjects', [])

    # 向后兼容：如果配置文件中没有 subjects 但有旧的 sheet_name 和 fill_rules，则转换为单学科
    if not subjects and 'sheet_name' in config and 'fill_rules' in config:
        subjects = [{
            "name": config.get('sheet_name', 'Sheet'),
            "sheet_name": config['sheet_name'],
            "fill_rules": config['fill_rules']
        }]

    if not input_dir or not output_dir or not subjects:
        print("错误：配置文件中必须包含 'input_dir', 'output_dir' 和 'subjects'（或向后兼容的 sheet_name+fill_rules）")
        return

    if not os.path.isdir(input_dir):
        print(f"错误：输入目录 {input_dir} 不存在")
        return

    xlsx_files = []
    for root, dirs, files in os.walk(input_dir):
        for file in files:
            if file.lower().endswith('.xlsx'):
                full_path = os.path.join(root, file)
                xlsx_files.append(full_path)

    if not xlsx_files:
        print(f"在 {input_dir} 中未找到任何 .xlsx 文件")
        return

    print(f"找到 {len(xlsx_files)} 个 Excel 文件，将处理 {len(subjects)} 个学科...")

    success_count = 0
    for input_path in xlsx_files:
        rel_path = os.path.relpath(input_path, input_dir)
        output_path = os.path.join(output_dir, rel_path)
        print(f"\n处理: {input_path} -> {output_path}")
        if process_one_excel(input_path, output_path, subjects):
            success_count += 1

    print(f"\n处理完成：成功 {success_count} / 总数 {len(xlsx_files)}")

if __name__ == '__main__':
    main()